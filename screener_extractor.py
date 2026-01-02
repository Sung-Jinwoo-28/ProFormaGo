import requests
import re
import io
from bs4 import BeautifulSoup
import pandas as pd
import sys
import os
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Color




class ScreenerExtractor:
    def __init__(self, ticker, years_limit=None, output_folder='.'):
        self.ticker = ticker.upper()
        self.years_limit = years_limit
        self.output_folder = output_folder
        self.base_url = f"https://www.screener.in/company/{self.ticker}/"
        self.soup = None
        self.data = {}

    def fetch_data(self):
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
        }
        
        # Try Consolidated URL first
        urls_to_try = [
            f"https://www.screener.in/company/{self.ticker}/consolidated/",
            f"https://www.screener.in/company/{self.ticker}/"
        ]
        
        for url in urls_to_try:
            print(f"Fetching data from: {url}")
            try:
                response = requests.get(url, headers=headers)
                if response.status_code == 200:
                    self.soup = BeautifulSoup(response.content, 'html.parser')
                    
                    # Basic validation - sometimes 404 pages return 200 with "Page not found" title
                    if "Page not found" in self.soup.title.text:
                        continue
                        
                    print("Successfully fetched data.")
                    return True
            except requests.exceptions.RequestException as e:
                print(f"Error fetching URL: {e}")
                continue
        
        print(f"Error: Could not fetch data for '{self.ticker}'. Checked consolidated and standalone pages.")
        return False

    def extract_table(self, section_id, table_name):
        if not self.soup:
            return

        section = self.soup.find('section', {'id': section_id})
        if not section:
            # Some sections might be named differently or missing
            # Try finding by text if ID fails, or just skip
            # specific logic for 'quarters' which sometimes is just a table w/o strict section ID wrap in some layouts?
            # actually screener is pretty consistent.
            print(f"Warning: Section '{section_id}' not found.")
            return

        table = section.find('table')
        if not table:
            print(f"Warning: Table not found in section '{section_id}'.")
            return

        # Extract headers
        headers = []
        thead = table.find('thead')
        if thead:
            header_row = thead.find('tr')
            header_row = thead.find('tr')
            headers = [th.get_text(strip=True) for th in header_row.find_all('th')]
            
            # Clean headers: "Mar 202415m" -> "Mar 2024"
            cleaned_headers = []
            for h in headers:
                # Match "Month Year" pattern
                match = re.match(r"([A-Za-z]{3} \d{4})", h)
                if match:
                    cleaned_headers.append(match.group(1))
                else:
                    cleaned_headers.append(h)
            headers = cleaned_headers
        
        # Extract rows
        rows = []
        tbody = table.find('tbody')
        if tbody:
            for tr in tbody.find_all('tr'):
                # Check if this is the "Raw PDF" row
                # We check the text of the first cell
                cells = tr.find_all(['td', 'th'])
                if not cells:
                    continue
                    
                first_cell_text = cells[0].get_text(strip=True)
                is_raw_pdf = "Raw PDF" in first_cell_text
                
                row_data = []
                for td in cells:
                    text = td.get_text(strip=True)
                    if is_raw_pdf:
                        link = td.find('a')
                        if link and 'href' in link.attrs:
                            href = link['href']
                            if href.startswith('/'):
                                href = f"https://www.screener.in{href}"
                            # Use Excel HYPERLINK formula
                            text = f'=HYPERLINK("{href}", "View File")'
                    row_data.append(text)
                
                if row_data:
                    rows.append(row_data)

        if not rows:
            return

        # Align headers and rows
        # Often the first column is empty in header but filled in rows (the metric name)
        if len(rows) > 0:
            cols_count = len(rows[0])
            if len(headers) < cols_count:
                headers = ['Metric'] + headers
            
            headers = headers[:cols_count]
            
            if len(headers) != cols_count:
                df = pd.DataFrame(rows)
            else:
                df = pd.DataFrame(rows, columns=headers)
            
            # Clean the data (convert to numbers)
            df = self.clean_df(df)
            
            # Filter columns by years if limit is set
            if self.years_limit and self.years_limit > 0:
                # Assuming structure: [Metric, Oldest Year, ..., Newest Year, (TTM?)]
                # We want to keep Metric + last N years
                
                # Identify 'Metric' column
                metric_col = df.columns[0]
                
                # Get the rest of the columns
                data_cols = df.columns[1:]
                
                # If TTM is present, keeps it? Usually yes.
                # Let's count TTM as one of the "data points" or separate?
                # User asked "how many eveys to data", likely meaning N years.
                # Safest bet: Keep TTM and last N actual years.
                # Screener columns are usually ascending.
                
                # Check for TTM
                has_ttm = 'TTM' in data_cols
                
                # Filter out TTM from year columns to count years properly
                year_cols = [c for c in data_cols if c != 'TTM']
                
                # Take the last N years
                selected_years = year_cols[-self.years_limit:]
                
                # Combine back: Metric + Selected Years + TTM (if available)
                # We excluded TTM from the 'limit' count, but we still want to show it.
                final_cols = [metric_col] + selected_years
                if has_ttm:
                    final_cols.append('TTM')
                
                # Filter, ensuring we don't crash if requesting more than available
                available_cols = [c for c in final_cols if c in df.columns]
                df = df[available_cols]

        else:
             return

        self.data[table_name] = df
        print(f"Extracted {table_name}")

    def clean_df(self, df):
        # Rename first column if it's empty or Unnamed
        if len(df.columns) > 0 and (df.columns[0] == '' or str(df.columns[0]).startswith('Unnamed')):
            df.rename(columns={df.columns[0]: 'Metric'}, inplace=True)

        for col in df.columns:
            # Skip the first column (Metric names)
            if col == 'Metric' or col == df.columns[0]:
                continue
            
            # Convert to numeric
            # We use apply to handle mixed types (numbers and percentages strings) in the same column
            if df[col].dtype == object:
                def convert_val(val):
                    if pd.isna(val) or val == '' or val == '--':
                        return None
                    
                    s_val = str(val).replace(',', '')
                    
                    # Check for percentage symbol
                    if '%' in s_val:
                        # Convert "15%" -> 0.15 to be a proper Excel percentage
                        try:
                            return float(s_val.replace('%', '')) / 100.0
                        except ValueError:
                            return s_val
                        
                    # Otherwise try to convert to number
                    try:
                        return float(s_val)
                    except ValueError:
                        return val
                
                df[col] = df[col].apply(convert_val)
        return df

    def run(self, in_memory=False):
        if not self.fetch_data():
            return None # Return None on failure

        # Sections to extract defined by their HTML ID
        sections = {
            'quarters': 'Quarterly Results',
            'profit-loss': 'Profit & Loss',
            'balance-sheet': 'Balance Sheet',
            'cash-flow': 'Cash Flows',
            'ratios': 'Ratios'
        }

        for sec_id, name in sections.items():
            self.extract_table(sec_id, name)
            
        self.extract_table('shareholding', 'Shareholding Pattern')
        
        # Calculate derived investing ratios
        self.calculate_investing_ratios()

        return self.save_to_excel(in_memory=in_memory)

    def extract_top_ratios(self):
        # Extract data from the top list (Market Cap, Current Price, High/Low, Stock P/E, etc.)
        if not self.soup:
            return {}
            
        top_ratios = {}
        ul = self.soup.find('ul', {'id': 'top-ratios'})
        if ul:
            for li in ul.find_all('li'):
                name_span = li.find('span', {'class': 'name'})
                value_span = li.find('span', {'class': 'number'})
                if name_span and value_span:
                    name = name_span.get_text(strip=True)
                    value = value_span.get_text(strip=True).replace(',', '')
                    try:
                        top_ratios[name] = float(value)
                    except ValueError:
                        top_ratios[name] = value
        return top_ratios

    def calculate_investing_ratios(self):
        if 'Profit & Loss' not in self.data or 'Balance Sheet' not in self.data:
            return

        print("Calculating Investing Ratios...")
        
        # Get Live/TTM Ratios
        top_ratios = self.extract_top_ratios()
        current_pe = top_ratios.get('Stock P/E')
        current_price = top_ratios.get('Current Price')

        # Helper to get numeric series safely
        def get_series(df, keyword):
            # Find row where Metric contains keyword (case insensitive)
            mask = df.iloc[:, 0].astype(str).str.contains(keyword, case=False, na=False)
            if mask.any():
                # Return the row as a series, excluding Metric column, properly index by columns
                row = df.loc[mask].iloc[0, 1:]
                return pd.to_numeric(row, errors='coerce')
            return None

        pnl = self.data['Profit & Loss']
        bs = self.data['Balance Sheet']
        
        # Identify common years
        common_cols = [c for c in pnl.columns if c in bs.columns and c != 'Metric' and c != 'TTM']
        
        # Check if TTM is available in P&L
        has_ttm = 'TTM' in pnl.columns
        
        # Helper to extend BS series with TTM (using last available year)
        def extend_bs_series(series):
            if series is None: return None
            if has_ttm and 'TTM' not in series.index and len(common_cols) > 0:
                # Use the last common year's value as proxy for TTM balance sheet items
                last_year = common_cols[-1]
                if last_year in series.index:
                    series['TTM'] = series[last_year]
            return series

        # Extract inputs & Extend BS items
        sales = get_series(pnl, 'Sales')
        if sales is None:
             sales = get_series(pnl, 'Revenue')

        net_profit = get_series(pnl, 'Net Profit')
        
        op_profit = get_series(pnl, 'Operating Profit')
        if op_profit is None:
             op_profit = get_series(pnl, 'Financing Profit')

        interest = get_series(pnl, 'Interest') # New
        eps = get_series(pnl, 'EPS')
        
        equity = extend_bs_series(get_series(bs, 'Equity Capital'))
        reserves = extend_bs_series(get_series(bs, 'Reserves'))
        borrowings = extend_bs_series(get_series(bs, 'Borrowings'))
        other_liab = extend_bs_series(get_series(bs, 'Other Liabilities'))
        total_assets = extend_bs_series(get_series(bs, 'Total Assets'))
        fixed_assets = extend_bs_series(get_series(bs, 'Fixed Assets'))
        cwip = extend_bs_series(get_series(bs, 'CWIP'))
        investments = extend_bs_series(get_series(bs, 'Investments'))
        
        shareholder_equity = None
        if equity is not None and reserves is not None:
            shareholder_equity = equity + reserves

        # Calculations
        metrics = {}
        
        # --- Profitability ---
        if net_profit is not None and sales is not None:
            metrics['Net Profit Margin %'] = (net_profit / sales) 
            
        if op_profit is not None and sales is not None:
             metrics['Operating Profit Margin %'] = (op_profit / sales)
             
        if sales is not None:
            metrics['Sales'] = sales
        
        if net_profit is not None and shareholder_equity is not None:
            metrics['Return on Equity (ROE) %'] = (net_profit / shareholder_equity)
            
        if op_profit is not None and shareholder_equity is not None and borrowings is not None:
             metrics['ROCE %'] = (op_profit / (shareholder_equity + borrowings))

        # --- Leverage ---
        if borrowings is not None and shareholder_equity is not None:
            metrics['Debt to Equity Ratio'] = (borrowings / shareholder_equity)
            
        if op_profit is not None and interest is not None:
             metrics['Interest Coverage Ratio'] = (op_profit / interest)

        if total_assets is not None and shareholder_equity is not None:
            metrics['Financial Leverage'] = (total_assets / shareholder_equity)

        # --- Efficiency ---
        if sales is not None and total_assets is not None:
            metrics['Asset Turnover Ratio'] = (sales / total_assets)

        # --- Liquidity (Approximated) ---
        # Current Assets approx = Total Assets - Fixed Assets - CWIP - Investments
        current_assets = None
        if total_assets is not None and fixed_assets is not None:
            current_assets = total_assets - fixed_assets
            if cwip is not None:
                current_assets = current_assets - cwip
            if investments is not None:
                current_assets = current_assets - investments
        
        # Current Liabilities approx = Other Liabilities (Trade Payables etc)
        current_liabs = other_liab
        
        if current_assets is not None and current_liabs is not None:
             metrics['Current Ratio'] = (current_assets / current_liabs)

        # --- Valuation (PE) ---
        # Construct DataFrame columns
        all_cols = common_cols + (['TTM'] if has_ttm else [])
        
        rows_list = []
        
        # Add PE Ratio manually
        pe_row = {'Metric': 'Price to Earning (PE)'}
        for col in all_cols:
            if col == 'TTM' and current_pe:
                 pe_row[col] = current_pe
            else:
                 pe_row[col] = None 
        rows_list.append(pe_row)

        for name, series in metrics.items():
            row = {'Metric': name}
            for col in all_cols:
                # series might have TTM now due to extend_bs_series or being from P&L
                if col in series.index:
                    row[col] = series[col]
                else:
                    row[col] = None
            rows_list.append(row)
            
        if rows_list:
            cols_order = ['Metric'] + all_cols
            ratios_df = pd.DataFrame(rows_list, columns=cols_order)
            self.data['Investing Ratios'] = ratios_df
            print("Created Investing Ratios sheet.")

    def save_to_excel(self, in_memory=False):
        if not self.data:
            print("No data extracted.")
            return None

        # If in_memory is True, we use BytesIO, otherwise filepath
        if in_memory:
             output = io.BytesIO()
             print("Generating file in memory...")
        else:
             filename = f"{self.ticker}_Financial_Model.xlsx"
             filepath = os.path.join(self.output_folder, filename)
             output = filepath
             print(f"Saving to {filepath}...")
        
        try:
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                for sheet_name, df in self.data.items():
                    # Excel sheet names max 31 chars
                    safe_sheet_name = sheet_name[:31]
                    df.to_excel(writer, sheet_name=safe_sheet_name, index=False)
                    
                    worksheet = writer.sheets[safe_sheet_name]
                    
                    # Auto-adjust column widths and apply formatting
                    for i, column_cells in enumerate(worksheet.columns):
                        # Adjust width
                        # Financial models usually have fixed width for data columns
                        # Column A (Metrics) gets more space
                        if i == 0:
                            worksheet.column_dimensions[column_cells[0].column_letter].width = 35
                        else:
                            worksheet.column_dimensions[column_cells[0].column_letter].width = 15
                    
                    # Define Styles
                    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid") # Financial Blue
                    header_font = Font(bold=True, color="FFFFFF")
                    thin_border = Border(left=Side(style='thin'), 
                                       right=Side(style='thin'), 
                                       top=Side(style='thin'), 
                                       bottom=Side(style='thin'))
                    center_align = Alignment(horizontal='center', vertical='center')
                    
                    # Apply Styles to Headers (Row 1)
                    for cell in worksheet[1]:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = center_align
                        cell.border = thin_border
                        
                    # Apply Styles to All Data Cells
                    for row in worksheet.iter_rows(min_row=2):
                        for cell in row:
                            cell.border = thin_border
                            
                    # Bold Metric Names (Column A)
                    for row in worksheet.iter_rows(min_row=2, max_col=1):
                        for cell in row:
                            cell.font = Font(bold=True)
                    
                    # Apply percentage formatting to rows where Metric name contains '%'
                    # OR if it's the Shareholding Pattern sheet (excluding 'No. of Shareholders')
                    for row in worksheet.iter_rows(min_row=2): # Skip header
                        metric_cell = row[0] # First column is Metric
                        
                        is_shareholding = (safe_sheet_name == 'Shareholding Pattern')
                        is_shareholder_count = (metric_cell.value and 'No. of Shareholders' in str(metric_cell.value))
                        should_format_percent = (metric_cell.value and '%' in str(metric_cell.value)) or (is_shareholding and not is_shareholder_count)

                        if should_format_percent:
                            for cell in row[1:]:
                                if isinstance(cell.value, (int, float)):
                                    cell.number_format = '0.00%'

            print("Done!")
            
            if in_memory:
                output.seek(0)
                return output
            else:
                print(f"File saved at: {os.path.abspath(output)}")
                return os.path.abspath(output)
        except Exception as e:
            print(f"Error saving Excel file: {e}")
            return None

if __name__ == "__main__":
    print("--- Screener.in Data Extractor ---")
    
    ticker = None
    years_limit = None

    if len(sys.argv) > 1:
        ticker = sys.argv[1]
        if len(sys.argv) > 2:
            try:
                years_limit = int(sys.argv[2])
            except ValueError:
                print("Invalid number for years. Use: python screener_extractor.py TICKER [YEARS]")
    else:
        ticker = input("Enter Company Ticker (e.g., RELIANCE, TCS): ").strip()
        years_input = input("How many years of data do you want? (Enter for all): ").strip()
        if years_input:
            try:
                years_limit = int(years_input)
            except ValueError:
                print("Invalid number, extracting all years.")
    
    if ticker:
        extractor = ScreenerExtractor(ticker, years_limit)
        extractor.run()
    else:
        print("No ticker provided. Exiting.")
